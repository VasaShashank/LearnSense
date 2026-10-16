"""
Multi-Format Source Materials Evaluation Benchmark.
Evaluates source fidelity (Concepts, Definitions, Relationships) across:
1. Microsoft Word (.docx)
2. Microsoft PowerPoint (.pptx)
3. Portable Document Format (.pdf)
4. Markdown / Plain Text (.md / .txt)
"""

import io
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from evaluation.phase2_eval import evaluate_knowledge_representation
from ingestion.document_formats import (
    docx_to_structured_document,
    pptx_to_structured_document,
    text_to_pdf,
)
from ingestion.pipeline import IngestionPipeline
from phase2.pipeline.runner import Phase2PipelineRunner


def _generate_golden_docx() -> bytes:
    doc = Document()
    doc.add_heading("Operating Systems: Memory Management", level=1)
    doc.add_paragraph("Virtual Memory is a memory management capability that provides an idealized abstraction of storage resources.")
    doc.add_paragraph("Paging is defined as a memory management scheme by which a computer stores and retrieves data from secondary storage.")
    doc.add_paragraph("Page Fault is an interrupt that occurs when a software program attempts to access a page not in RAM.")
    doc.add_paragraph("Virtual Memory requires Paging as an underlying architectural mechanism.")
    doc.add_paragraph("Paging requires Page Table to translate virtual addresses to physical frames.")

    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Concept"
    table.cell(0, 1).text = "Access Latency"
    table.cell(1, 0).text = "RAM"
    table.cell(1, 1).text = "100 ns"
    table.cell(2, 0).text = "Disk Storage"
    table.cell(2, 1).text = "10 ms"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _generate_golden_pptx() -> bytes:
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Distributed Systems
    s1 = prs.slides.add_slide(blank_layout)
    tx_box = s1.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    tx_box.text_frame.text = "Distributed Systems Architecture"

    body_box = s1.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4))
    body_box.text_frame.text = (
        "Consensus Algorithm is defined as a fault-tolerant protocol for reaching agreement among distributed nodes.\n"
        "Raft Consensus is a consensus algorithm designed to be understandable.\n"
        "Raft Consensus is an example of Consensus Algorithm.\n"
        "Distributed Systems require Consensus Algorithm for fault tolerance."
    )
    s1.notes_slide.notes_text_frame.text = "Lecturer Note: Emphasize that Raft decomposes consensus into leader election and log replication."

    # Slide 2: Table
    s2 = prs.slides.add_slide(blank_layout)
    s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(2))
    table = s2.shapes[-1].table
    table.cell(0, 0).text = "Protocol"
    table.cell(0, 1).text = "Complexity"
    table.cell(1, 0).text = "Paxos"
    table.cell(1, 1).text = "High"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _generate_golden_md() -> str:
    return (
        "# Machine Learning Fundamentals\n\n"
        "Gradient Descent: An optimization algorithm used to train machine learning models by minimizing errors.\n\n"
        "Loss Function: A mathematical method of evaluating how well an algorithm models the given dataset.\n\n"
        "Neural Network: A computational learning system inspired by biological neural networks.\n\n"
        "Deep Learning is a subset of Machine Learning.\n\n"
        "Deep Learning requires Gradient Descent to optimize high-dimensional neural network weights.\n\n"
        "Backpropagation requires Calculus Derivatives as a computational foundation."
    )


def test_docx_source_fidelity_benchmark():
    docx_bytes = _generate_golden_docx()
    struct_doc = docx_to_structured_document(docx_bytes, "os_memory.docx", "doc_os_bench")

    runner = Phase2PipelineRunner()
    ekr = runner.process(struct_doc)

    golden = {
        "concepts": ["Virtual Memory", "Paging", "Page Fault"],
        "definitions": {
            "virtual memory": "memory management capability",
            "paging": "memory management scheme",
            "page fault": "interrupt that occurs",
        },
        "relationships": [
            {"source": "virtual memory", "target": "paging"},
            {"source": "paging", "target": "page table"},
        ],
    }

    metrics = evaluate_knowledge_representation(ekr, golden)
    assert metrics["concept_recall"] >= 0.66
    assert metrics["definition_recall"] >= 0.66
    assert metrics["total_extracted_concepts"] >= 3


def test_pptx_source_fidelity_benchmark():
    pptx_bytes = _generate_golden_pptx()
    struct_doc = pptx_to_structured_document(pptx_bytes, "distributed_systems.pptx", "doc_ds_bench")

    runner = Phase2PipelineRunner()
    ekr = runner.process(struct_doc)

    golden = {
        "concepts": ["Consensus Algorithm", "Raft Consensus", "Distributed Systems"],
        "definitions": {
            "consensus algorithm": "fault-tolerant protocol",
        },
        "relationships": [
            {"source": "distributed systems", "target": "consensus algorithm"},
        ],
    }

    metrics = evaluate_knowledge_representation(ekr, golden)
    assert metrics["concept_recall"] >= 0.66
    assert metrics["definition_recall"] == 1.0


def test_markdown_and_colon_source_fidelity_benchmark():
    md_text = _generate_golden_md()
    pdf_bytes = text_to_pdf(md_text)
    pipeline = IngestionPipeline()
    struct_doc = pipeline.process_pdf_bytes(pdf_bytes, "ml_guide.md", "doc_ml_bench")

    runner = Phase2PipelineRunner()
    ekr = runner.process(struct_doc)

    golden = {
        "concepts": ["Gradient Descent", "Loss Function", "Neural Network", "Deep Learning"],
        "definitions": {
            "gradient descent": "optimization algorithm",
            "loss function": "mathematical method",
            "neural network": "computational learning system",
        },
        "relationships": [
            {"source": "deep learning", "target": "gradient descent"},
        ],
    }

    metrics = evaluate_knowledge_representation(ekr, golden)
    assert metrics["concept_recall"] >= 0.5
    assert metrics["definition_recall"] >= 0.66
