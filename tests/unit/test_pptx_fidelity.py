"""
Unit tests for PPTX extraction fidelity and Stage 4/5 definition binding.
Verifies Section 3.1 & 5.1 of LearnSense Source Fidelity requirements.
"""

import io
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from ingestion.document_formats import pptx_to_structured_document
from schemas.document import BlockTypeEnum
from phase2.pipeline.stage1_normalize import run_stage1_normalize
from phase2.pipeline.stage2_context import run_stage2_context
from phase2.pipeline.stage3_segment import run_stage3_segmentation
from phase2.pipeline.stage4_candidates import extract_candidates
from phase2.pipeline.stage5_resolution import run_stage5_entity_resolution
from phase2.models import EvidenceKindEnum


def _create_sample_pptx_bytes() -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 1. Title shape
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    tf = title_box.text_frame
    tf.text = "Distributed Systems Architecture"

    # 2. Text bullet
    body_box = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(8), Inches(1))
    p = body_box.text_frame.paragraphs[0]
    p.text = "Microservices are autonomous software services deployed independently."

    # 3. Table shape (2 rows, 2 columns)
    rows, cols = 2, 2
    table_shape = slide.shapes.add_table(rows, cols, Inches(1), Inches(3.2), Inches(6), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Architecture Model"
    table.cell(0, 1).text = "Scalability"
    table.cell(1, 0).text = "Monolith"
    table.cell(1, 1).text = "Vertical"

    # 4. Speaker notes
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = "Remind students that microservices trade architectural simplicity for operational complexity."

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_pptx_fidelity_table_and_notes():
    raw_bytes = _create_sample_pptx_bytes()
    doc = pptx_to_structured_document(raw_bytes, "architecture_lecture.pptx", "doc_test_arch")

    assert doc.document_id == "doc_test_arch"
    assert len(doc.pages) == 1

    page = doc.pages[0]
    assert page.width == 720.0  # 10 inches * 72 pt/in
    assert page.height == 540.0  # 7.5 inches * 72 pt/in

    block_types = [b.type for b in page.blocks]
    block_roles = [b.role for b in page.blocks]

    # Verify heading
    assert BlockTypeEnum.HEADING in block_types
    assert "heading_l1" in block_roles

    # Verify table extracted with TABLE type and markdown representation
    assert BlockTypeEnum.TABLE in block_types
    table_block = next(b for b in page.blocks if b.type == BlockTypeEnum.TABLE)
    assert table_block.role == "table"
    assert "| Architecture Model | Scalability |" in table_block.content.text
    assert "| Monolith | Vertical |" in table_block.content.text

    # Verify speaker notes extracted
    notes_block = next((b for b in page.blocks if b.role == "speaker_notes"), None)
    assert notes_block is not None
    assert "operational complexity" in notes_block.content.text


def test_definition_retention_and_binding():
    # Build a document with a single-occurrence defined term
    sample_text = (
        "Quantum Supremacy is defined as the potential ability of quantum devices to solve problems intractable for classical computers.\n"
        "We discuss classical computing benchmarks here.\n"
        "Modern supercomputers perform exaflop calculations daily."
    )
    from ingestion.document_formats import text_to_pdf
    from ingestion.pipeline import IngestionPipeline

    pdf_bytes = text_to_pdf(sample_text)
    pipeline = IngestionPipeline()
    struct_doc = pipeline.process_pdf_bytes(pdf_bytes, "quantum.txt", "doc_quantum_01")

    from phase2.adapters.nlp_adapters import Tier01DeterministicAdapter

    # Run Phase 2 stages
    norm_doc = run_stage1_normalize(struct_doc)
    doc_ctx = run_stage2_context(norm_doc)
    units = run_stage3_segmentation(norm_doc, doc_ctx, Tier01DeterministicAdapter())
    candidates = extract_candidates(norm_doc, doc_ctx, units)

    # 1. Verify Quantum Supremacy was harvested as definition
    harvested = [d for d in candidates.harvested_definitions if "quantum supremacy" in d["term"].lower()]
    assert len(harvested) > 0, "Definition pattern must harvest Quantum Supremacy"

    # 2. Run Stage 5 resolution
    concepts, _ = run_stage5_entity_resolution(norm_doc, candidates)
    qs_concept = next((c for c in concepts if "quantum supremacy" in c.canonical_name.lower()), None)

    assert qs_concept is not None, "Single-occurrence defined term must be preserved in concepts"
    assert qs_concept.description is not None
    assert "potential ability of quantum devices" in qs_concept.description

    # 3. Verify definition evidence kind is present
    def_ev = [ev for ev in candidates.evidence if ev.kind == EvidenceKindEnum.DEFINITION_PATTERN]
    assert len(def_ev) > 0
    assert any(ev.evidence_id in qs_concept.evidence_ids for ev in def_ev)
