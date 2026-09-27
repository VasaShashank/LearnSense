"""
Document Input Validator for Taproot Phase 5.
Validates file type, size, magic header, page limits, and corrupt/empty files prior to ingestion.
Supports PDF (.pdf), PowerPoint (.pptx), Word (.docx), and Images (.png, .jpg, .jpeg).
"""

from typing import Optional
from pathlib import Path
import io
import pymupdf as fitz
from PIL import Image
from phase5.config.phase5_config import phase5_config
from phase5.errors.error_types import InputValidationError
from phase5.models.validation_result import RecoveryClassification, ValidationResult, ValidationStatus
from phase5.observability.validation_events import ValidationEventLogger


class InputValidator:
    """Validates raw incoming file uploads prior to resource-intensive processing."""

    SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".docx", ".png", ".jpg", ".jpeg"}

    def __init__(
        self,
        max_file_size_bytes: int = phase5_config.MAX_FILE_SIZE_BYTES,
        max_page_count: int = phase5_config.MAX_PAGE_COUNT,
    ):
        self.max_file_size_bytes = max_file_size_bytes
        self.max_page_count = max_page_count

    def validate_file(self, content: bytes, filename: str) -> ValidationResult:
        result = ValidationResult(status=ValidationStatus.VALID, is_valid=True)
        size_bytes = len(content)

        # 1. File type extension check
        ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
        if ext not in self.SUPPORTED_EXTENSIONS:
            result.add_error(f"Unsupported file type for filename '{filename}'. Supported types: PDF, PPTX, DOCX, PNG, JPG, JPEG.")
            result.recoverable = False
            result.recovery_classification = RecoveryClassification.NON_RECOVERABLE
            ValidationEventLogger.log_event("input_validation_failed", "INVALID", result.errors[-1])
            return result

        # 2. File size check
        if size_bytes > self.max_file_size_bytes:
            result.add_error(f"File size ({size_bytes} bytes) exceeds configured limit of {self.max_file_size_bytes} bytes.")
            result.recoverable = False
            result.recovery_classification = RecoveryClassification.NON_RECOVERABLE
            ValidationEventLogger.log_event("input_validation_failed", "INVALID", result.errors[-1])
            return result

        if size_bytes < 10:
            result.add_error("Uploaded file is empty or corrupted (file size too small).")
            result.recoverable = False
            result.recovery_classification = RecoveryClassification.NON_RECOVERABLE
            ValidationEventLogger.log_event("input_validation_failed", "INVALID", result.errors[-1])
            return result

        # 3. Format-specific validation & structure checks
        if ext == ".pdf":
            if not content.startswith(b"%PDF-"):
                offset = content.find(b"%PDF-")
                if offset == -1:
                    result.add_error("Missing PDF magic header '%PDF-'. File is not a valid PDF.")
                    result.recoverable = False
                    result.recovery_classification = RecoveryClassification.NON_RECOVERABLE
                    ValidationEventLogger.log_event("input_validation_failed", "INVALID", result.errors[-1])
                    return result

            try:
                doc = fitz.open(stream=content, filetype="pdf")
                if doc.is_encrypted and doc.needs_pass:
                    doc.close()
                    result.add_error("PDF is password-protected and requires a password to decrypt.")
                    result.recoverable = False
                    result.recovery_classification = RecoveryClassification.NON_RECOVERABLE
                    ValidationEventLogger.log_event("input_validation_failed", "INVALID", result.errors[-1])
                    return result

                page_count = len(doc)
                doc.close()

                if page_count == 0:
                    result.add_error("PDF document contains zero pages.")
                    result.recoverable = False
                    result.recovery_classification = RecoveryClassification.NON_RECOVERABLE
                    ValidationEventLogger.log_event("input_validation_failed", "INVALID", result.errors[-1])
                    return result

                if page_count > self.max_page_count:
                    result.add_error(f"Document page count ({page_count}) exceeds limit of {self.max_page_count} pages.")
                    result.recoverable = False
                    result.recovery_classification = RecoveryClassification.NON_RECOVERABLE
                    ValidationEventLogger.log_event("input_validation_failed", "INVALID", result.errors[-1])
                    return result

                result.metadata["page_count"] = page_count
                result.metadata["size_bytes"] = size_bytes

            except Exception as e:
                result.add_error(f"Failed to parse PDF document structure: {str(e)}", action="Attempt PDF repair.")
                result.recoverable = True
                result.recovery_classification = RecoveryClassification.RECOVERABLE
                ValidationEventLogger.log_event("input_validation_failed", "RECOVERABLE", result.errors[-1])
                return result

        elif ext == ".docx":
            try:
                import docx
                doc = docx.Document(io.BytesIO(content))
                page_count = max(1, len(doc.paragraphs) // 5)
                result.metadata["page_count"] = page_count
                result.metadata["size_bytes"] = size_bytes
            except Exception as e:
                result.add_error(f"Failed to parse DOCX document structure: {str(e)}")
                result.recoverable = False
                result.recovery_classification = RecoveryClassification.NON_RECOVERABLE
                ValidationEventLogger.log_event("input_validation_failed", "INVALID", result.errors[-1])
                return result

        elif ext == ".pptx":
            try:
                import pptx
                prs = pptx.Presentation(io.BytesIO(content))
                page_count = len(prs.slides)
                if page_count == 0:
                    result.add_error("PPTX presentation contains zero slides.")
                    result.recoverable = False
                    result.recovery_classification = RecoveryClassification.NON_RECOVERABLE
                    ValidationEventLogger.log_event("input_validation_failed", "INVALID", result.errors[-1])
                    return result
                result.metadata["page_count"] = page_count
                result.metadata["size_bytes"] = size_bytes
            except Exception as e:
                result.add_error(f"Failed to parse PPTX presentation structure: {str(e)}")
                result.recoverable = False
                result.recovery_classification = RecoveryClassification.NON_RECOVERABLE
                ValidationEventLogger.log_event("input_validation_failed", "INVALID", result.errors[-1])
                return result

        elif ext in (".png", ".jpg", ".jpeg"):
            try:
                img = Image.open(io.BytesIO(content))
                img.verify()
                result.metadata["page_count"] = 1
                result.metadata["size_bytes"] = size_bytes
            except Exception as e:
                result.add_error(f"Failed to parse image file: {str(e)}")
                result.recoverable = False
                result.recovery_classification = RecoveryClassification.NON_RECOVERABLE
                ValidationEventLogger.log_event("input_validation_failed", "INVALID", result.errors[-1])
                return result

        ValidationEventLogger.log_event("input_validation_passed", "VALID", f"File '{filename}' passed input validation.")
        return result
